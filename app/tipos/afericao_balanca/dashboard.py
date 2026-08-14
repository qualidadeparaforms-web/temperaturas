from datetime import date, datetime, time, timedelta
from io import BytesIO

from flask import Blueprint, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.timezone_utils import agora_brasilia
from app.tipos.afericao_balanca.models import Balanca, RegistroAfericaoBalanca

dashboard_bp = Blueprint("afericao_balanca_dashboard", __name__, url_prefix="/dashboard/afericao_balanca")

COR_CABECALHO = "1F3864"
COR_CABECALHO_GRUPO = "2E75B6"
COR_FUNDO_FORA_PADRAO = "FFC7CE"
COR_FONTE_FORA_PADRAO = "C00000"


def periodo_para_datas(periodo: str, data_inicio: str | None = None, data_fim: str | None = None):
    hoje = agora_brasilia().date()
    if periodo == "dia":
        return hoje, hoje
    if periodo == "semana":
        return hoje - timedelta(days=6), hoje
    if periodo == "mes":
        return hoje - timedelta(days=29), hoje
    if periodo == "personalizado" and data_inicio and data_fim:
        try:
            return (
                datetime.strptime(data_inicio, "%Y-%m-%d").date(),
                datetime.strptime(data_fim, "%Y-%m-%d").date(),
            )
        except ValueError:
            pass
    return hoje - timedelta(days=6), hoje


def buscar_registros(periodo: str, balanca: str, data_inicio: str | None = None, data_fim: str | None = None):
    inicio, fim = periodo_para_datas(periodo, data_inicio, data_fim)
    query = RegistroAfericaoBalanca.query.filter(
        RegistroAfericaoBalanca.data >= inicio, RegistroAfericaoBalanca.data <= fim
    )
    if balanca and balanca != "todos":
        try:
            query = query.filter(RegistroAfericaoBalanca.balanca_id == int(balanca))
        except ValueError:
            pass
    registros = query.order_by(RegistroAfericaoBalanca.data.asc(), RegistroAfericaoBalanca.id.asc()).all()
    return registros, inicio, fim


def _parametros_filtro():
    return {
        "periodo": request.args.get("periodo", "mes"),
        "balanca": request.args.get("balanca", "todos"),
        "data_inicio": request.args.get("data_inicio"),
        "data_fim": request.args.get("data_fim"),
    }


def _balancas_por_setor() -> dict:
    balancas = Balanca.query.order_by(Balanca.id.asc()).all()
    agrupadas: dict[str, list] = {}
    for balanca in balancas:
        agrupadas.setdefault(balanca.setor, []).append(balanca)
    return agrupadas


@dashboard_bp.route("")
def index():
    from app.tipos import TIPOS_REGISTRO

    return render_template(
        "tipos/afericao_balanca/dashboard.html",
        balancas_por_setor=_balancas_por_setor(),
        tipos=TIPOS_REGISTRO,
    )


@dashboard_bp.route("/api")
def api_registros():
    filtros = _parametros_filtro()
    registros, inicio, fim = buscar_registros(**filtros)

    total = len(registros)
    fora_padrao = sum(1 for r in registros if not r.conforme)
    conformes = total - fora_padrao

    return jsonify(
        {
            "registros": [r.to_dict() for r in registros],
            "resumo": {
                "total": total,
                "conformes": conformes,
                "fora_padrao": fora_padrao,
                "percentual_conforme": round(conformes / total * 100, 1) if total else 100.0,
            },
            "periodo": {
                "inicio": inicio.strftime("%d/%m/%Y"),
                "fim": fim.strftime("%d/%m/%Y"),
            },
        }
    )


def _estilizar_cabecalho(ws, num_colunas: int, cor: str = COR_CABECALHO) -> None:
    preenchimento = PatternFill(start_color=cor, end_color=cor, fill_type="solid")
    fonte = Font(color="FFFFFF", bold=True)
    for indice in range(1, num_colunas + 1):
        celula = ws.cell(row=1, column=indice)
        celula.fill = preenchimento
        celula.font = fonte
        celula.alignment = Alignment(horizontal="center")


def adicionar_planilha(wb: Workbook, inicio: date, fim: date, filtros_extra: dict | None = None) -> None:
    """Acrescenta ao workbook uma aba no formato de matriz original:
    as 20 balanças nas linhas (sempre todas, agrupadas por setor, cada
    grupo com sua própria linha de cabeçalho), dias do período nas
    colunas, cada célula com a leitura (g) e o status C/NC. Ignora o
    filtro "balanca" do painel de propósito — a matriz sempre mostra
    todas as balanças."""
    registros = (
        RegistroAfericaoBalanca.query.filter(
            RegistroAfericaoBalanca.data >= inicio, RegistroAfericaoBalanca.data <= fim
        )
        .order_by(RegistroAfericaoBalanca.data.asc(), RegistroAfericaoBalanca.id.asc())
        .all()
    )

    dias = sorted({r.data for r in registros})

    # Se houver mais de uma leitura da mesma balança no mesmo dia
    # (ex.: correção), a mais recente prevalece na matriz.
    registro_por_balanca_dia: dict[int, dict] = {}
    for registro in registros:
        registro_por_balanca_dia.setdefault(registro.balanca_id, {})[registro.data] = registro

    ws = wb.create_sheet("Aferição Balanças")

    cabecalho = ["Balança"] + [d.strftime("%d/%m") for d in dias]
    ws.append(cabecalho)
    _estilizar_cabecalho(ws, len(cabecalho))

    preenchimento_grupo = PatternFill(start_color=COR_CABECALHO_GRUPO, end_color=COR_CABECALHO_GRUPO, fill_type="solid")
    fonte_grupo = Font(color="FFFFFF", bold=True)
    preenchimento_fora_padrao = PatternFill(
        start_color=COR_FUNDO_FORA_PADRAO, end_color=COR_FUNDO_FORA_PADRAO, fill_type="solid"
    )
    fonte_fora_padrao = Font(color=COR_FONTE_FORA_PADRAO, bold=True)

    for setor, balancas_do_setor in _balancas_por_setor().items():
        ws.append([setor] + ["" for _ in dias])
        linha_grupo = ws.max_row
        for coluna in range(1, len(dias) + 2):
            celula = ws.cell(row=linha_grupo, column=coluna)
            celula.fill = preenchimento_grupo
            celula.font = fonte_grupo
        if dias:
            ws.merge_cells(start_row=linha_grupo, start_column=1, end_row=linha_grupo, end_column=len(dias) + 1)

        for balanca in balancas_do_setor:
            registros_da_balanca = registro_por_balanca_dia.get(balanca.id, {})
            linha_valores = [f"Equip. {balanca.equipamento}"]
            for d in dias:
                registro = registros_da_balanca.get(d)
                linha_valores.append(f"{registro.leitura:.1f}g ({registro.status})" if registro else "")
            ws.append(linha_valores)
            linha = ws.max_row
            for coluna, d in enumerate(dias, start=2):
                registro = registros_da_balanca.get(d)
                if registro and registro.status == "NC":
                    celula = ws.cell(row=linha, column=coluna)
                    celula.fill = preenchimento_fora_padrao
                    celula.font = fonte_fora_padrao
            ws.cell(row=linha, column=1).alignment = Alignment(horizontal="left")

    ws.column_dimensions["A"].width = 45
    for coluna in range(2, len(dias) + 2):
        ws.column_dimensions[get_column_letter(coluna)].width = 14


@dashboard_bp.route("/exportar")
def exportar():
    filtros = _parametros_filtro()
    inicio, fim = periodo_para_datas(filtros["periodo"], filtros["data_inicio"], filtros["data_fim"])

    wb = Workbook()
    del wb["Sheet"]
    adicionar_planilha(wb, inicio, fim, {})

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nome_arquivo = f"registros_afericao_balanca_{inicio.strftime('%Y%m%d')}_{fim.strftime('%Y%m%d')}.xlsx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=nome_arquivo,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def contar(inicio: date, fim: date) -> int:
    return RegistroAfericaoBalanca.query.filter(
        RegistroAfericaoBalanca.data >= inicio, RegistroAfericaoBalanca.data <= fim
    ).count()


def linhas_combinadas(inicio: date, fim: date) -> list:
    registros = (
        RegistroAfericaoBalanca.query.filter(
            RegistroAfericaoBalanca.data >= inicio, RegistroAfericaoBalanca.data <= fim
        )
        .order_by(RegistroAfericaoBalanca.data.asc(), RegistroAfericaoBalanca.id.asc())
        .all()
    )
    return [
        {
            "tipo": "PAC 08-G - Aferição das Balanças",
            "icone": "⚙️",
            "data": r.data.strftime("%d/%m/%Y"),
            "horario": "",
            "resumo": f"{r.balanca.rotulo() if r.balanca else '?'}: {r.leitura:.1f}g "
            f"({'Conforme' if r.status == 'C' else '⚠️ NC'})",
            "conforme": r.conforme,
            "ordenacao": datetime.combine(r.data, time.min),
        }
        for r in registros
    ]
