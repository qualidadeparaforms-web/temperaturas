from datetime import date, datetime, timedelta
from io import BytesIO

from flask import Blueprint, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.timezone_utils import agora_brasilia
from app.tipos.integridade_componente.models import (
    EQUIPAMENTOS,
    RegistroIntegridadeComponente,
    VerificacaoRT,
)

dashboard_bp = Blueprint(
    "integridade_componente_dashboard", __name__, url_prefix="/dashboard/integridade_componente"
)

COR_CABECALHO = "1F3864"
COR_FUNDO_FORA_PADRAO = "FFC7CE"
COR_FONTE_FORA_PADRAO = "C00000"


def periodo_para_datas(periodo: str, data_inicio: str | None = None, data_fim: str | None = None):
    hoje = agora_brasilia().date()
    if periodo == "dia":
        return hoje, hoje
    if periodo == "semana":
        return hoje - timedelta(days=6), hoje
    if periodo == "mes":
        return hoje - timedelta(days=29), hoje
    if periodo == "personalizado" and data_inicio and data_fim:
        try:
            return (
                datetime.strptime(data_inicio, "%Y-%m-%d").date(),
                datetime.strptime(data_fim, "%Y-%m-%d").date(),
            )
        except ValueError:
            pass
    return hoje - timedelta(days=6), hoje


def buscar_registros(
    periodo: str, equipamento: str, data_inicio: str | None = None, data_fim: str | None = None
):
    inicio, fim = periodo_para_datas(periodo, data_inicio, data_fim)
    query = RegistroIntegridadeComponente.query.filter(
        RegistroIntegridadeComponente.data >= inicio, RegistroIntegridadeComponente.data <= fim
    )
    if equipamento and equipamento != "todos":
        query = query.filter(RegistroIntegridadeComponente.equipamento == equipamento)
    registros = query.order_by(
        RegistroIntegridadeComponente.data.asc(), RegistroIntegridadeComponente.horario.asc()
    ).all()
    return registros, inicio, fim


def _parametros_filtro():
    return {
        "periodo": request.args.get("periodo", "semana"),
        "equipamento": request.args.get("equipamento", "todos"),
        "data_inicio": request.args.get("data_inicio"),
        "data_fim": request.args.get("data_fim"),
    }


@dashboard_bp.route("")
def index():
    from app.tipos import TIPOS_REGISTRO

    return render_template(
        "tipos/integridade_componente/dashboard.html", equipamentos=EQUIPAMENTOS, tipos=TIPOS_REGISTRO
    )


@dashboard_bp.route("/api")
def api_registros():
    filtros = _parametros_filtro()
    registros, inicio, fim = buscar_registros(**filtros)

    total = len(registros)
    fora_padrao = sum(1 for r in registros if not r.conforme)
    conformes = total - fora_padrao

    return jsonify(
        {
            "registros": [r.to_dict() for r in registros],
            "resumo": {
                "total": total,
                "conformes": conformes,
                "fora_padrao": fora_padrao,
                "percentual_conforme": round(conformes / total * 100, 1) if total else 100.0,
            },
            "periodo": {
                "inicio": inicio.strftime("%d/%m/%Y"),
                "fim": fim.strftime("%d/%m/%Y"),
            },
        }
    )


def _estilizar_cabecalho(ws, num_colunas: int) -> None:
    cabecalho_preenchimento = PatternFill(start_color=COR_CABECALHO, end_color=COR_CABECALHO, fill_type="solid")
    cabecalho_fonte = Font(color="FFFFFF", bold=True)
    for indice in range(1, num_colunas + 1):
        celula = ws.cell(row=1, column=indice)
        celula.fill = cabecalho_preenchimento
        celula.font = cabecalho_fonte
        celula.alignment = Alignment(horizontal="center")


def adicionar_planilha(wb: Workbook, inicio: date, fim: date, filtros_extra: dict | None = None) -> None:
    """Acrescenta ao workbook DUAS abas: os registros de integridade
    (checklist C/NC) e o histórico de verificações RT do período —
    tabelas independentes, sem relação pai/filho entre si."""
    filtros_extra = filtros_extra or {}
    equipamento = filtros_extra.get("equipamento", "todos")

    query = RegistroIntegridadeComponente.query.filter(
        RegistroIntegridadeComponente.data >= inicio, RegistroIntegridadeComponente.data <= fim
    )
    if equipamento and equipamento != "todos":
        query = query.filter(RegistroIntegridadeComponente.equipamento == equipamento)
    registros = query.order_by(
        RegistroIntegridadeComponente.data.asc(), RegistroIntegridadeComponente.horario.asc()
    ).all()

    preenchimento_fora_padrao = PatternFill(
        start_color=COR_FUNDO_FORA_PADRAO, end_color=COR_FUNDO_FORA_PADRAO, fill_type="solid"
    )
    fonte_fora_padrao = Font(color=COR_FONTE_FORA_PADRAO, bold=True)

    # Aba 1: registros de integridade (checklist C/NC).
    ws = wb.create_sheet("Integridade Componentes")
    colunas = ["Dia", "Horário", "Equipamento", "Componente", "Momento", "Responsável", "Status"]
    ws.append(colunas)
    _estilizar_cabecalho(ws, len(colunas))

    for registro in registros:
        ws.append(
            [
                registro.data.strftime("%d/%m/%Y"),
                registro.horario.strftime("%H:%M"),
                registro.equipamento,
                registro.componente,
                registro.momento,
                registro.responsavel,
                registro.status,
            ]
        )
        if not registro.conforme:
            linha = ws.max_row
            for indice in range(1, len(colunas) + 1):
                celula = ws.cell(row=linha, column=indice)
                celula.fill = preenchimento_fora_padrao
                celula.font = fonte_fora_padrao

    for indice, largura in enumerate([12, 10, 24, 20, 18, 24, 10], start=1):
        ws.column_dimensions[chr(64 + indice)].width = largura

    # Aba 2: histórico de verificações RT (conferências rápidas, sem
    # status de conformidade) — sempre com todos os equipamentos,
    # ignora o filtro de equipamento do painel de propósito.
    verificacoes = (
        VerificacaoRT.query.filter(
            VerificacaoRT.data_hora >= datetime.combine(inicio, datetime.min.time()),
            VerificacaoRT.data_hora < datetime.combine(fim + timedelta(days=1), datetime.min.time()),
        )
        .order_by(VerificacaoRT.data_hora.asc())
        .all()
    )

    ws_rt = wb.create_sheet("Verificações RT")
    colunas_rt = ["Dia", "Horário", "Equipamento"]
    ws_rt.append(colunas_rt)
    _estilizar_cabecalho(ws_rt, len(colunas_rt))
    for verificacao in verificacoes:
        ws_rt.append(
            [
                verificacao.data_hora.strftime("%d/%m/%Y"),
                verificacao.data_hora.strftime("%H:%M"),
                verificacao.equipamento,
            ]
        )
    for indice, largura in enumerate([12, 10, 24], start=1):
        ws_rt.column_dimensions[chr(64 + indice)].width = largura


@dashboard_bp.route("/exportar")
def exportar():
    filtros = _parametros_filtro()
    inicio, fim = periodo_para_datas(filtros["periodo"], filtros["data_inicio"], filtros["data_fim"])

    wb = Workbook()
    del wb["Sheet"]
    adicionar_planilha(wb, inicio, fim, {"equipamento": filtros["equipamento"]})

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nome_arquivo = (
        f"registros_integridade_componente_{inicio.strftime('%Y%m%d')}_{fim.strftime('%Y%m%d')}.xlsx"
    )
    return send_file(
        buffer,
        as_attachment=True,
        download_name=nome_arquivo,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def contar(inicio: date, fim: date) -> int:
    return RegistroIntegridadeComponente.query.filter(
        RegistroIntegridadeComponente.data >= inicio, RegistroIntegridadeComponente.data <= fim
    ).count()


def linhas_combinadas(inicio: date, fim: date) -> list:
    registros = (
        RegistroIntegridadeComponente.query.filter(
            RegistroIntegridadeComponente.data >= inicio, RegistroIntegridadeComponente.data <= fim
        )
        .order_by(RegistroIntegridadeComponente.data.asc(), RegistroIntegridadeComponente.horario.asc())
        .all()
    )
    return [
        {
            "tipo": "PAC 17 - Integridade de Componentes",
            "icone": "🪡",
            "data": r.data.strftime("%d/%m/%Y"),
            "horario": r.horario.strftime("%H:%M"),
            "resumo": f"{r.equipamento} — {r.componente} ({r.momento}): {'Conforme' if r.status == 'C' else '⚠️ NC'}",
            "conforme": r.conforme,
            "ordenacao": datetime.combine(r.data, r.horario),
        }
        for r in registros
    ]
