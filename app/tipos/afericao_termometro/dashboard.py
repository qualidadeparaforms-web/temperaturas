from datetime import date, datetime, time, timedelta
from io import BytesIO

from flask import Blueprint, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.timezone_utils import agora_brasilia
from app.tipos.afericao_termometro.models import (
    DIFERENCA_MAXIMA_C,
    RegistroAfericaoTermometro,
    Termometro,
)

dashboard_bp = Blueprint("afericao_termometro_dashboard", __name__, url_prefix="/dashboard/afericao_termometro")

COR_CABECALHO = "1F3864"
COR_FUNDO_FORA_PADRAO = "FFC7CE"
COR_FONTE_FORA_PADRAO = "C00000"


def periodo_para_datas(periodo: str, data_inicio: str | None = None, data_fim: str | None = None):
    hoje = agora_brasilia().date()
    if periodo == "dia":
        return hoje, hoje
    if periodo == "semana":
        return hoje - timedelta(days=6), hoje
    if periodo == "mes":
        return hoje - timedelta(days=29), hoje
    if periodo == "personalizado" and data_inicio and data_fim:
        try:
            return (
                datetime.strptime(data_inicio, "%Y-%m-%d").date(),
                datetime.strptime(data_fim, "%Y-%m-%d").date(),
            )
        except ValueError:
            pass
    return hoje - timedelta(days=6), hoje


def buscar_registros(periodo: str, termometro: str, data_inicio: str | None = None, data_fim: str | None = None):
    inicio, fim = periodo_para_datas(periodo, data_inicio, data_fim)
    query = RegistroAfericaoTermometro.query.filter(
        RegistroAfericaoTermometro.data >= inicio, RegistroAfericaoTermometro.data <= fim
    )
    if termometro and termometro != "todos":
        try:
            query = query.filter(RegistroAfericaoTermometro.termometro_id == int(termometro))
        except ValueError:
            pass
    registros = query.order_by(RegistroAfericaoTermometro.data.asc(), RegistroAfericaoTermometro.id.asc()).all()
    return registros, inicio, fim


def _parametros_filtro():
    return {
        "periodo": request.args.get("periodo", "mes"),
        "termometro": request.args.get("termometro", "todos"),
        "data_inicio": request.args.get("data_inicio"),
        "data_fim": request.args.get("data_fim"),
    }


@dashboard_bp.route("")
def index():
    from app.tipos import TIPOS_REGISTRO

    termometros = Termometro.query.order_by(Termometro.id.asc()).all()
    return render_template(
        "tipos/afericao_termometro/dashboard.html",
        termometros=termometros,
        tipos=TIPOS_REGISTRO,
    )


@dashboard_bp.route("/api")
def api_registros():
    filtros = _parametros_filtro()
    registros, inicio, fim = buscar_registros(**filtros)

    total = len(registros)
    fora_padrao = sum(1 for r in registros if not r.conforme)
    conformes = total - fora_padrao

    return jsonify(
        {
            "registros": [r.to_dict() for r in registros],
            "resumo": {
                "total": total,
                "conformes": conformes,
                "fora_padrao": fora_padrao,
                "percentual_conforme": round(conformes / total * 100, 1) if total else 100.0,
            },
            "periodo": {
                "inicio": inicio.strftime("%d/%m/%Y"),
                "fim": fim.strftime("%d/%m/%Y"),
            },
        }
    )


def adicionar_planilha(wb: Workbook, inicio: date, fim: date, filtros_extra: dict | None = None) -> None:
    """Acrescenta ao workbook uma aba no formato da planilha original:
    uma linha por aferição (não uma matriz pivotada como PAC 08-E/G) —
    data, nº do equipamento, equipamento, as 4 leituras de
    temperatura, variação aceitável e C/NC, responsável. Ignora o
    filtro "termometro" do painel de propósito — a exportação sempre
    mostra todas as aferições do período."""
    registros = (
        RegistroAfericaoTermometro.query.filter(
            RegistroAfericaoTermometro.data >= inicio, RegistroAfericaoTermometro.data <= fim
        )
        .order_by(RegistroAfericaoTermometro.data.asc(), RegistroAfericaoTermometro.id.asc())
        .all()
    )

    ws = wb.create_sheet("Aferição Termômetros")

    colunas = [
        "Dia", "Nº Equipamento", "Equipamento",
        "Temp. Quente Padrão (°C)", "Temp. Quente Equipamento (°C)",
        "Temp. Fria Padrão (°C)", "Temp. Fria Equipamento (°C)",
        "Variação Aceitável (°C)", "C/NC", "Responsável",
    ]
    ws.append(colunas)

    cabecalho_preenchimento = PatternFill(start_color=COR_CABECALHO, end_color=COR_CABECALHO, fill_type="solid")
    cabecalho_fonte = Font(color="FFFFFF", bold=True)
    for indice in range(1, len(colunas) + 1):
        celula = ws.cell(row=1, column=indice)
        celula.fill = cabecalho_preenchimento
        celula.font = cabecalho_fonte
        celula.alignment = Alignment(horizontal="center")

    preenchimento_fora_padrao = PatternFill(
        start_color=COR_FUNDO_FORA_PADRAO, end_color=COR_FUNDO_FORA_PADRAO, fill_type="solid"
    )
    fonte_fora_padrao = Font(color=COR_FONTE_FORA_PADRAO, bold=True)

    for registro in registros:
        termometro = registro.termometro
        ws.append(
            [
                registro.data.strftime("%d/%m/%Y"),
                termometro.codigo if termometro else "",
                termometro.descricao if termometro else "",
                registro.temp_quente_padrao,
                registro.temp_quente_equipamento,
                registro.temp_fria_padrao,
                registro.temp_fria_equipamento,
                f"±{DIFERENCA_MAXIMA_C:g}",
                registro.status,
                registro.responsavel,
            ]
        )
        if not registro.conforme:
            linha = ws.max_row
            for indice in range(1, len(colunas) + 1):
                celula = ws.cell(row=linha, column=indice)
                celula.fill = preenchimento_fora_padrao
                celula.font = fonte_fora_padrao

    larguras = [12, 16, 30, 14, 16, 14, 16, 12, 8, 20]
    for indice, largura in enumerate(larguras, start=1):
        ws.column_dimensions[chr(64 + indice)].width = largura


@dashboard_bp.route("/exportar")
def exportar():
    filtros = _parametros_filtro()
    inicio, fim = periodo_para_datas(filtros["periodo"], filtros["data_inicio"], filtros["data_fim"])

    wb = Workbook()
    del wb["Sheet"]
    adicionar_planilha(wb, inicio, fim, {})

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nome_arquivo = f"registros_afericao_termometro_{inicio.strftime('%Y%m%d')}_{fim.strftime('%Y%m%d')}.xlsx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=nome_arquivo,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def contar(inicio: date, fim: date) -> int:
    return RegistroAfericaoTermometro.query.filter(
        RegistroAfericaoTermometro.data >= inicio, RegistroAfericaoTermometro.data <= fim
    ).count()


def linhas_combinadas(inicio: date, fim: date) -> list:
    registros = (
        RegistroAfericaoTermometro.query.filter(
            RegistroAfericaoTermometro.data >= inicio, RegistroAfericaoTermometro.data <= fim
        )
        .order_by(RegistroAfericaoTermometro.data.asc(), RegistroAfericaoTermometro.id.asc())
        .all()
    )
    return [
        {
            "tipo": "PAC 08-F - Aferição dos Termômetros",
            "icone": "🎯",
            "data": r.data.strftime("%d/%m/%Y"),
            "horario": "",
            "resumo": f"{r.termometro.rotulo() if r.termometro else '?'}: "
            f"diferença {max(r.diferenca_quente, r.diferenca_fria):.1f}°C "
            f"({'Conforme' if r.status == 'C' else '⚠️ NC'})",
            "conforme": r.conforme,
            "ordenacao": datetime.combine(r.data, time.min),
        }
        for r in registros
    ]
