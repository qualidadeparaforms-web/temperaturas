from datetime import date, datetime, time, timedelta
from io import BytesIO

from flask import Blueprint, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.timezone_utils import agora_brasilia
from app.tipos.iluminacao.models import PONTOS_POR_SETOR, RegistroIluminacao

dashboard_bp = Blueprint("iluminacao_dashboard", __name__, url_prefix="/dashboard/iluminacao")

COR_CABECALHO = "1F3864"
COR_FUNDO_FORA_PADRAO = "FFC7CE"
COR_FONTE_FORA_PADRAO = "C00000"


def periodo_para_datas(periodo: str, data_inicio: str | None = None, data_fim: str | None = None):
    hoje = agora_brasilia().date()
    if periodo == "dia":
        return hoje, hoje
    if periodo == "semana":
        return hoje - timedelta(days=6), hoje
    if periodo == "mes":
        return hoje - timedelta(days=29), hoje
    if periodo == "personalizado" and data_inicio and data_fim:
        try:
            return (
                datetime.strptime(data_inicio, "%Y-%m-%d").date(),
                datetime.strptime(data_fim, "%Y-%m-%d").date(),
            )
        except ValueError:
            pass
    return hoje - timedelta(days=6), hoje


def buscar_registros(periodo: str, ponto: str, data_inicio: str | None = None, data_fim: str | None = None):
    inicio, fim = periodo_para_datas(periodo, data_inicio, data_fim)
    query = RegistroIluminacao.query.filter(
        RegistroIluminacao.data >= inicio, RegistroIluminacao.data <= fim
    )
    # "ponto" chega como "<setor>|<local>" — o par é a identidade do
    # ponto, já que "local" sozinho pode repetir entre setores (ex.:
    # "Barreira sanitária" existe em 2 setores diferentes).
    if ponto and ponto != "todos" and "|" in ponto:
        setor, local = ponto.split("|", 1)
        query = query.filter(RegistroIluminacao.setor == setor, RegistroIluminacao.local == local)
    registros = query.order_by(RegistroIluminacao.data.asc(), RegistroIluminacao.id.asc()).all()
    return registros, inicio, fim


def _parametros_filtro():
    return {
        "periodo": request.args.get("periodo", "mes"),
        "ponto": request.args.get("ponto", "todos"),
        "data_inicio": request.args.get("data_inicio"),
        "data_fim": request.args.get("data_fim"),
    }


@dashboard_bp.route("")
def index():
    from app.tipos import TIPOS_REGISTRO

    return render_template(
        "tipos/iluminacao/dashboard.html",
        pontos_por_setor=PONTOS_POR_SETOR,
        tipos=TIPOS_REGISTRO,
    )


@dashboard_bp.route("/api")
def api_registros():
    filtros = _parametros_filtro()
    registros, inicio, fim = buscar_registros(**filtros)

    total = len(registros)
    fora_padrao = sum(1 for r in registros if not r.conforme)
    conformes = total - fora_padrao

    return jsonify(
        {
            "registros": [r.to_dict() for r in registros],
            "resumo": {
                "total": total,
                "conformes": conformes,
                "fora_padrao": fora_padrao,
                "percentual_conforme": round(conformes / total * 100, 1) if total else 100.0,
            },
            "periodo": {
                "inicio": inicio.strftime("%d/%m/%Y"),
                "fim": fim.strftime("%d/%m/%Y"),
            },
        }
    )


def adicionar_planilha(wb: Workbook, inicio: date, fim: date, filtros_extra: dict | None = None) -> None:
    """Acrescenta ao workbook uma aba no formato original da planilha:
    uma linha por leitura (não uma matriz pivotada) — dia, setor,
    local, lux necessário, lux obtido, C/NC. Ignora o filtro "ponto"
    do painel de propósito — a exportação sempre mostra todas as
    leituras do período."""
    registros = (
        RegistroIluminacao.query.filter(RegistroIluminacao.data >= inicio, RegistroIluminacao.data <= fim)
        .order_by(RegistroIluminacao.data.asc(), RegistroIluminacao.id.asc())
        .all()
    )

    ws = wb.create_sheet("Iluminação")

    colunas = ["Dia", "Setor", "Local", "Lux Necessário", "Lux Obtido", "C/NC"]
    ws.append(colunas)

    cabecalho_preenchimento = PatternFill(start_color=COR_CABECALHO, end_color=COR_CABECALHO, fill_type="solid")
    cabecalho_fonte = Font(color="FFFFFF", bold=True)
    for indice in range(1, len(colunas) + 1):
        celula = ws.cell(row=1, column=indice)
        celula.fill = cabecalho_preenchimento
        celula.font = cabecalho_fonte
        celula.alignment = Alignment(horizontal="center")

    preenchimento_fora_padrao = PatternFill(
        start_color=COR_FUNDO_FORA_PADRAO, end_color=COR_FUNDO_FORA_PADRAO, fill_type="solid"
    )
    fonte_fora_padrao = Font(color=COR_FONTE_FORA_PADRAO, bold=True)

    for registro in registros:
        ws.append(
            [
                registro.data.strftime("%d/%m/%Y"),
                registro.setor,
                registro.local,
                registro.lux_necessario,
                registro.lux_obtido,
                registro.status,
            ]
        )
        if not registro.conforme:
            linha = ws.max_row
            for indice in range(1, len(colunas) + 1):
                celula = ws.cell(row=linha, column=indice)
                celula.fill = preenchimento_fora_padrao
                celula.font = fonte_fora_padrao

    larguras = [12, 26, 32, 16, 14, 8]
    for indice, largura in enumerate(larguras, start=1):
        ws.column_dimensions[chr(64 + indice)].width = largura


@dashboard_bp.route("/exportar")
def exportar():
    filtros = _parametros_filtro()
    inicio, fim = periodo_para_datas(filtros["periodo"], filtros["data_inicio"], filtros["data_fim"])

    wb = Workbook()
    del wb["Sheet"]
    adicionar_planilha(wb, inicio, fim, {})

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nome_arquivo = f"registros_iluminacao_{inicio.strftime('%Y%m%d')}_{fim.strftime('%Y%m%d')}.xlsx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=nome_arquivo,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def contar(inicio: date, fim: date) -> int:
    return RegistroIluminacao.query.filter(
        RegistroIluminacao.data >= inicio, RegistroIluminacao.data <= fim
    ).count()


def linhas_combinadas(inicio: date, fim: date) -> list:
    registros = (
        RegistroIluminacao.query.filter(RegistroIluminacao.data >= inicio, RegistroIluminacao.data <= fim)
        .order_by(RegistroIluminacao.data.asc(), RegistroIluminacao.id.asc())
        .all()
    )
    return [
        {
            "tipo": "PAC 08-D - Monitoramento da Iluminação",
            "icone": "💡",
            "data": r.data.strftime("%d/%m/%Y"),
            "horario": "",
            "resumo": f"{r.setor} — {r.local}: {r.lux_obtido:g} lux "
            f"(mín. {r.lux_necessario}) — {'Conforme' if r.status == 'C' else '⚠️ NC'}",
            "conforme": r.conforme,
            "ordenacao": datetime.combine(r.data, time.min),
        }
        for r in registros
    ]
