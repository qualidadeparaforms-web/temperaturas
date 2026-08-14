from datetime import date, datetime, time, timedelta
from io import BytesIO

from flask import Blueprint, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.timezone_utils import agora_brasilia
from app.tipos.ventilacao.models import SETORES, SETORES_POR_GRUPO, RegistroVentilacao

dashboard_bp = Blueprint("ventilacao_dashboard", __name__, url_prefix="/dashboard/ventilacao")

COR_CABECALHO = "1F3864"
COR_CABECALHO_GRUPO = "2E75B6"
COR_FUNDO_FORA_PADRAO = "FFC7CE"
COR_FONTE_FORA_PADRAO = "C00000"


def periodo_para_datas(periodo: str, data_inicio: str | None = None, data_fim: str | None = None):
    hoje = agora_brasilia().date()
    if periodo == "dia":
        return hoje, hoje
    if periodo == "semana":
        return hoje - timedelta(days=6), hoje
    if periodo == "mes":
        return hoje - timedelta(days=29), hoje
    if periodo == "personalizado" and data_inicio and data_fim:
        try:
            return (
                datetime.strptime(data_inicio, "%Y-%m-%d").date(),
                datetime.strptime(data_fim, "%Y-%m-%d").date(),
            )
        except ValueError:
            pass
    return hoje - timedelta(days=6), hoje


def buscar_registros(periodo: str, setor: str, data_inicio: str | None = None, data_fim: str | None = None):
    inicio, fim = periodo_para_datas(periodo, data_inicio, data_fim)
    query = RegistroVentilacao.query.filter(
        RegistroVentilacao.data >= inicio, RegistroVentilacao.data <= fim
    )
    if setor and setor != "todos":
        query = query.filter(RegistroVentilacao.setor == setor)
    registros = query.order_by(RegistroVentilacao.data.asc(), RegistroVentilacao.id.asc()).all()
    return registros, inicio, fim


def _parametros_filtro():
    return {
        "periodo": request.args.get("periodo", "mes"),
        "setor": request.args.get("setor", "todos"),
        "data_inicio": request.args.get("data_inicio"),
        "data_fim": request.args.get("data_fim"),
    }


@dashboard_bp.route("")
def index():
    from app.tipos import TIPOS_REGISTRO

    return render_template(
        "tipos/ventilacao/dashboard.html",
        setores_por_grupo=SETORES_POR_GRUPO,
        tipos=TIPOS_REGISTRO,
    )


@dashboard_bp.route("/api")
def api_registros():
    filtros = _parametros_filtro()
    registros, inicio, fim = buscar_registros(**filtros)

    total = len(registros)
    fora_padrao = sum(1 for r in registros if not r.conforme)
    conformes = total - fora_padrao

    return jsonify(
        {
            "registros": [r.to_dict() for r in registros],
            "resumo": {
                "total": total,
                "conformes": conformes,
                "fora_padrao": fora_padrao,
                "percentual_conforme": round(conformes / total * 100, 1) if total else 100.0,
            },
            "periodo": {
                "inicio": inicio.strftime("%d/%m/%Y"),
                "fim": fim.strftime("%d/%m/%Y"),
            },
        }
    )


def _estilizar_cabecalho(ws, num_colunas: int, cor: str = COR_CABECALHO) -> None:
    preenchimento = PatternFill(start_color=cor, end_color=cor, fill_type="solid")
    fonte = Font(color="FFFFFF", bold=True)
    for indice in range(1, num_colunas + 1):
        celula = ws.cell(row=1, column=indice)
        celula.fill = preenchimento
        celula.font = fonte
        celula.alignment = Alignment(horizontal="center")


def adicionar_planilha(wb: Workbook, inicio: date, fim: date, filtros_extra: dict | None = None) -> None:
    """Acrescenta ao workbook uma aba no formato de matriz original:
    os 22 setores nas linhas (sempre todos, na ordem fixa, agrupados
    por "Setor Produção" e "Setor Embalagem Secundária e Expedição",
    cada grupo com uma linha de cabeçalho própria), dias do período
    nas colunas, C/NC em cada célula. Ignora o filtro "setor" do
    painel de propósito — a matriz sempre mostra todos os setores."""
    registros = (
        RegistroVentilacao.query.filter(RegistroVentilacao.data >= inicio, RegistroVentilacao.data <= fim)
        .order_by(RegistroVentilacao.data.asc(), RegistroVentilacao.id.asc())
        .all()
    )

    dias = sorted({r.data for r in registros})

    # Se houver mais de um registro do mesmo setor no mesmo dia (ex.:
    # correção), o último (por ordem de inserção) prevalece na matriz.
    status_por_setor_dia: dict[str, dict] = {setor: {} for setor in SETORES}
    for registro in registros:
        status_por_setor_dia.setdefault(registro.setor, {})[registro.data] = registro.status

    ws = wb.create_sheet("Ventilação")

    cabecalho = ["Setor"] + [d.strftime("%d/%m") for d in dias]
    ws.append(cabecalho)
    _estilizar_cabecalho(ws, len(cabecalho))

    preenchimento_grupo = PatternFill(start_color=COR_CABECALHO_GRUPO, end_color=COR_CABECALHO_GRUPO, fill_type="solid")
    fonte_grupo = Font(color="FFFFFF", bold=True)
    preenchimento_fora_padrao = PatternFill(
        start_color=COR_FUNDO_FORA_PADRAO, end_color=COR_FUNDO_FORA_PADRAO, fill_type="solid"
    )
    fonte_fora_padrao = Font(color=COR_FONTE_FORA_PADRAO, bold=True)

    for grupo, setores_do_grupo in SETORES_POR_GRUPO.items():
        ws.append([grupo] + ["" for _ in dias])
        linha_grupo = ws.max_row
        for coluna in range(1, len(dias) + 2):
            celula = ws.cell(row=linha_grupo, column=coluna)
            celula.fill = preenchimento_grupo
            celula.font = fonte_grupo
        if dias:
            ws.merge_cells(start_row=linha_grupo, start_column=1, end_row=linha_grupo, end_column=len(dias) + 1)

        for setor in setores_do_grupo:
            status_do_setor = status_por_setor_dia.get(setor, {})
            ws.append([setor] + [status_do_setor.get(d, "") for d in dias])
            linha = ws.max_row
            for coluna, d in enumerate(dias, start=2):
                if status_do_setor.get(d) == "NC":
                    celula = ws.cell(row=linha, column=coluna)
                    celula.fill = preenchimento_fora_padrao
                    celula.font = fonte_fora_padrao
            ws.cell(row=linha, column=1).alignment = Alignment(horizontal="left")

    ws.column_dimensions["A"].width = 50
    for coluna in range(2, len(dias) + 2):
        ws.column_dimensions[get_column_letter(coluna)].width = 8


@dashboard_bp.route("/exportar")
def exportar():
    filtros = _parametros_filtro()
    inicio, fim = periodo_para_datas(filtros["periodo"], filtros["data_inicio"], filtros["data_fim"])

    wb = Workbook()
    del wb["Sheet"]
    adicionar_planilha(wb, inicio, fim, {})

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nome_arquivo = f"registros_ventilacao_{inicio.strftime('%Y%m%d')}_{fim.strftime('%Y%m%d')}.xlsx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=nome_arquivo,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def contar(inicio: date, fim: date) -> int:
    return RegistroVentilacao.query.filter(
        RegistroVentilacao.data >= inicio, RegistroVentilacao.data <= fim
    ).count()


def linhas_combinadas(inicio: date, fim: date) -> list:
    registros = (
        RegistroVentilacao.query.filter(RegistroVentilacao.data >= inicio, RegistroVentilacao.data <= fim)
        .order_by(RegistroVentilacao.data.asc(), RegistroVentilacao.id.asc())
        .all()
    )
    return [
        {
            "tipo": "PAC 08-E - Monitoramento da Ventilação",
            "icone": "🌬️",
            "data": r.data.strftime("%d/%m/%Y"),
            "horario": "",
            "resumo": f"{r.setor}: {'Conforme' if r.status == 'C' else '⚠️ NC'}",
            "conforme": r.conforme,
            "ordenacao": datetime.combine(r.data, time.min),
        }
        for r in registros
    ]
