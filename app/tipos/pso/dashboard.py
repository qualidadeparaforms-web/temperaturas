from datetime import date, datetime, time, timedelta
from io import BytesIO

from flask import Blueprint, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.timezone_utils import agora_brasilia
from app.tipos.pso.models import PSOS, RegistroPso

dashboard_bp = Blueprint("pso_dashboard", __name__, url_prefix="/dashboard/pso")

COR_CABECALHO = "1F3864"
COR_FUNDO_FORA_PADRAO = "FFC7CE"
COR_FONTE_FORA_PADRAO = "C00000"


def periodo_para_datas(periodo: str, data_inicio: str | None = None, data_fim: str | None = None):
    hoje = agora_brasilia().date()
    if periodo == "dia":
        return hoje, hoje
    if periodo == "semana":
        return hoje - timedelta(days=6), hoje
    if periodo == "mes":
        return hoje - timedelta(days=29), hoje
    if periodo == "personalizado" and data_inicio and data_fim:
        try:
            return (
                datetime.strptime(data_inicio, "%Y-%m-%d").date(),
                datetime.strptime(data_fim, "%Y-%m-%d").date(),
            )
        except ValueError:
            pass
    return hoje - timedelta(days=6), hoje


def buscar_registros(periodo: str, pso: str, data_inicio: str | None = None, data_fim: str | None = None):
    inicio, fim = periodo_para_datas(periodo, data_inicio, data_fim)
    query = RegistroPso.query.filter(RegistroPso.data >= inicio, RegistroPso.data <= fim)
    if pso and pso != "todos":
        query = query.filter(RegistroPso.pso == pso)
    registros = query.order_by(RegistroPso.data.asc(), RegistroPso.id.asc()).all()
    return registros, inicio, fim


def _parametros_filtro():
    return {
        "periodo": request.args.get("periodo", "semana"),
        "pso": request.args.get("pso", "todos"),
        "data_inicio": request.args.get("data_inicio"),
        "data_fim": request.args.get("data_fim"),
    }


@dashboard_bp.route("")
def index():
    from app.tipos import TIPOS_REGISTRO

    return render_template("tipos/pso/dashboard.html", psos=PSOS, tipos=TIPOS_REGISTRO)


@dashboard_bp.route("/api")
def api_registros():
    filtros = _parametros_filtro()
    registros, inicio, fim = buscar_registros(**filtros)

    total = len(registros)
    fora_padrao = sum(1 for r in registros if not r.conforme)
    conformes = total - fora_padrao

    return jsonify(
        {
            "registros": [r.to_dict() for r in registros],
            "resumo": {
                "total": total,
                "conformes": conformes,
                "fora_padrao": fora_padrao,
                "percentual_conforme": round(conformes / total * 100, 1) if total else 100.0,
            },
            "periodo": {
                "inicio": inicio.strftime("%d/%m/%Y"),
                "fim": fim.strftime("%d/%m/%Y"),
            },
        }
    )


def _estilizar_cabecalho(ws, num_colunas: int) -> None:
    cabecalho_preenchimento = PatternFill(start_color=COR_CABECALHO, end_color=COR_CABECALHO, fill_type="solid")
    cabecalho_fonte = Font(color="FFFFFF", bold=True)
    for indice in range(1, num_colunas + 1):
        celula = ws.cell(row=1, column=indice)
        celula.fill = cabecalho_preenchimento
        celula.font = cabecalho_fonte
        celula.alignment = Alignment(horizontal="center")


def adicionar_planilha(wb: Workbook, inicio: date, fim: date, filtros_extra: dict | None = None) -> None:
    """Acrescenta ao workbook uma aba no formato de matriz original:
    PSOs nas linhas (sempre os 7, na ordem fixa), dias do período nas
    colunas, C/NC em cada célula. Ignora o filtro "pso" do painel de
    propósito — a matriz sempre mostra todos os PSOs."""
    registros = (
        RegistroPso.query.filter(RegistroPso.data >= inicio, RegistroPso.data <= fim)
        .order_by(RegistroPso.data.asc(), RegistroPso.id.asc())
        .all()
    )

    dias = sorted({r.data for r in registros})

    # Se houver mais de um registro do mesmo PSO no mesmo dia (ex.:
    # correção), o último (por ordem de inserção) prevalece na matriz.
    status_por_pso_dia: dict[str, dict] = {pso: {} for pso in PSOS}
    for registro in registros:
        status_por_pso_dia.setdefault(registro.pso, {})[registro.data] = registro.status

    ws = wb.create_sheet("PSOs")

    cabecalho = ["PSO"] + [d.strftime("%d/%m") for d in dias]
    ws.append(cabecalho)
    _estilizar_cabecalho(ws, len(cabecalho))

    preenchimento_fora_padrao = PatternFill(
        start_color=COR_FUNDO_FORA_PADRAO, end_color=COR_FUNDO_FORA_PADRAO, fill_type="solid"
    )
    fonte_fora_padrao = Font(color=COR_FONTE_FORA_PADRAO, bold=True)

    for pso in PSOS:
        status_do_pso = status_por_pso_dia.get(pso, {})
        ws.append([pso] + [status_do_pso.get(d, "") for d in dias])
        linha = ws.max_row
        for coluna, d in enumerate(dias, start=2):
            if status_do_pso.get(d) == "NC":
                celula = ws.cell(row=linha, column=coluna)
                celula.fill = preenchimento_fora_padrao
                celula.font = fonte_fora_padrao
        ws.cell(row=linha, column=1).alignment = Alignment(horizontal="left")

    ws.column_dimensions["A"].width = 50
    for coluna in range(2, len(dias) + 2):
        ws.column_dimensions[get_column_letter(coluna)].width = 8


@dashboard_bp.route("/exportar")
def exportar():
    filtros = _parametros_filtro()
    inicio, fim = periodo_para_datas(filtros["periodo"], filtros["data_inicio"], filtros["data_fim"])

    wb = Workbook()
    del wb["Sheet"]
    adicionar_planilha(wb, inicio, fim, {})

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nome_arquivo = f"registros_pso_{inicio.strftime('%Y%m%d')}_{fim.strftime('%Y%m%d')}.xlsx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=nome_arquivo,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def contar(inicio: date, fim: date) -> int:
    return RegistroPso.query.filter(RegistroPso.data >= inicio, RegistroPso.data <= fim).count()


def linhas_combinadas(inicio: date, fim: date) -> list:
    registros = (
        RegistroPso.query.filter(RegistroPso.data >= inicio, RegistroPso.data <= fim)
        .order_by(RegistroPso.data.asc(), RegistroPso.id.asc())
        .all()
    )
    return [
        {
            "tipo": "PAC 11 - Monitoramento dos PSO's",
            "icone": "🧼",
            "data": r.data.strftime("%d/%m/%Y"),
            "horario": "",
            "resumo": f"{r.pso}: {'Conforme' if r.status == 'C' else '⚠️ NC'}",
            "conforme": r.conforme,
            "ordenacao": datetime.combine(r.data, time.min),
        }
        for r in registros
    ]
