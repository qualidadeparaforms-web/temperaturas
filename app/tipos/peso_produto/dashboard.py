from datetime import date, datetime, timedelta
from io import BytesIO

from flask import Blueprint, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.timezone_utils import agora_brasilia
from app.tipos.peso_produto.models import RegistroPesoProduto

dashboard_bp = Blueprint("peso_produto_dashboard", __name__, url_prefix="/dashboard/peso_produto")

COR_CABECALHO = "1F3864"
COR_FUNDO_FORA_PADRAO = "FFC7CE"
COR_FONTE_FORA_PADRAO = "C00000"


def periodo_para_datas(periodo: str, data_inicio: str | None = None, data_fim: str | None = None):
    hoje = agora_brasilia().date()
    if periodo == "dia":
        return hoje, hoje
    if periodo == "semana":
        return hoje - timedelta(days=6), hoje
    if periodo == "mes":
        return hoje - timedelta(days=29), hoje
    if periodo == "personalizado" and data_inicio and data_fim:
        try:
            return (
                datetime.strptime(data_inicio, "%Y-%m-%d").date(),
                datetime.strptime(data_fim, "%Y-%m-%d").date(),
            )
        except ValueError:
            pass
    return hoje - timedelta(days=6), hoje


def buscar_registros(periodo: str, data_inicio: str | None = None, data_fim: str | None = None):
    inicio, fim = periodo_para_datas(periodo, data_inicio, data_fim)
    registros = (
        RegistroPesoProduto.query.filter(
            RegistroPesoProduto.data >= inicio, RegistroPesoProduto.data <= fim
        )
        .order_by(RegistroPesoProduto.data.asc(), RegistroPesoProduto.horario.asc())
        .all()
    )
    return registros, inicio, fim


def _parametros_filtro():
    return {
        "periodo": request.args.get("periodo", "semana"),
        "data_inicio": request.args.get("data_inicio"),
        "data_fim": request.args.get("data_fim"),
    }


@dashboard_bp.route("")
def index():
    from app.tipos import TIPOS_REGISTRO

    return render_template("tipos/peso_produto/dashboard.html", tipos=TIPOS_REGISTRO)


@dashboard_bp.route("/api")
def api_registros():
    filtros = _parametros_filtro()
    registros, inicio, fim = buscar_registros(**filtros)

    total_registros = len(registros)
    total_pesagens = sum(r.total_pesagens for r in registros)
    total_nc = sum(r.total_nc for r in registros)
    percentual_conforme = (
        round((total_pesagens - total_nc) / total_pesagens * 100, 1) if total_pesagens else 100.0
    )

    return jsonify(
        {
            "registros": [r.to_dict() for r in registros],
            "resumo": {
                "total_registros": total_registros,
                "total_pesagens": total_pesagens,
                "total_nc": total_nc,
                "percentual_conforme": percentual_conforme,
            },
            "periodo": {
                "inicio": inicio.strftime("%d/%m/%Y"),
                "fim": fim.strftime("%d/%m/%Y"),
            },
        }
    )


def _estilizar_cabecalho(ws, num_colunas: int) -> None:
    cabecalho_preenchimento = PatternFill(start_color=COR_CABECALHO, end_color=COR_CABECALHO, fill_type="solid")
    cabecalho_fonte = Font(color="FFFFFF", bold=True)
    for indice in range(1, num_colunas + 1):
        celula = ws.cell(row=1, column=indice)
        celula.fill = cabecalho_preenchimento
        celula.font = cabecalho_fonte
        celula.alignment = Alignment(horizontal="center")


def adicionar_planilha(wb: Workbook, inicio: date, fim: date, filtros_extra: dict | None = None) -> None:
    """Acrescenta ao workbook DUAS abas: um resumo por produto/dia
    (uma linha por registro) e o detalhe de cada pesagem individual."""
    registros = (
        RegistroPesoProduto.query.filter(
            RegistroPesoProduto.data >= inicio, RegistroPesoProduto.data <= fim
        )
        .order_by(RegistroPesoProduto.data.asc(), RegistroPesoProduto.horario.asc())
        .all()
    )

    preenchimento_fora_padrao = PatternFill(
        start_color=COR_FUNDO_FORA_PADRAO, end_color=COR_FUNDO_FORA_PADRAO, fill_type="solid"
    )
    fonte_fora_padrao = Font(color=COR_FONTE_FORA_PADRAO, bold=True)

    # Aba 1: resumo por produto/dia.
    ws_resumo = wb.create_sheet("Peso Produto")
    colunas_resumo = [
        "Dia",
        "Produto",
        "Peso Líquido Nominal",
        "Peso Embalagem",
        "Total Pesagens",
        "Total NC",
        "Responsável",
    ]
    ws_resumo.append(colunas_resumo)
    _estilizar_cabecalho(ws_resumo, len(colunas_resumo))

    for registro in registros:
        ws_resumo.append(
            [
                registro.data.strftime("%d/%m/%Y"),
                registro.produto,
                registro.peso_liquido_nominal,
                registro.peso_embalagem,
                registro.total_pesagens,
                registro.total_nc,
                registro.responsavel,
            ]
        )
        if registro.total_nc > 0:
            linha = ws_resumo.max_row
            for indice in range(1, len(colunas_resumo) + 1):
                celula = ws_resumo.cell(row=linha, column=indice)
                celula.fill = preenchimento_fora_padrao
                celula.font = fonte_fora_padrao

    for indice, largura in enumerate([12, 24, 18, 16, 14, 10, 24], start=1):
        ws_resumo.column_dimensions[chr(64 + indice)].width = largura

    # Aba 2: detalhe de cada pesagem individual.
    ws_detalhe = wb.create_sheet("Peso Produto Detalhe")
    colunas_detalhe = ["Dia", "Horário", "Produto", "Peso Medido", "Conforme"]
    ws_detalhe.append(colunas_detalhe)
    _estilizar_cabecalho(ws_detalhe, len(colunas_detalhe))

    for registro in registros:
        for pesagem in registro.pesagens:
            ws_detalhe.append(
                [
                    registro.data.strftime("%d/%m/%Y"),
                    registro.horario.strftime("%H:%M"),
                    registro.produto,
                    pesagem.peso_medido,
                    "Sim" if pesagem.conforme else "Não",
                ]
            )
            if not pesagem.conforme:
                linha = ws_detalhe.max_row
                for indice in range(1, len(colunas_detalhe) + 1):
                    celula = ws_detalhe.cell(row=linha, column=indice)
                    celula.fill = preenchimento_fora_padrao
                    celula.font = fonte_fora_padrao

    for indice, largura in enumerate([12, 10, 24, 14, 12], start=1):
        ws_detalhe.column_dimensions[chr(64 + indice)].width = largura


@dashboard_bp.route("/exportar")
def exportar():
    filtros = _parametros_filtro()
    inicio, fim = periodo_para_datas(filtros["periodo"], filtros["data_inicio"], filtros["data_fim"])

    wb = Workbook()
    del wb["Sheet"]
    adicionar_planilha(wb, inicio, fim, {})

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nome_arquivo = f"registros_peso_produto_{inicio.strftime('%Y%m%d')}_{fim.strftime('%Y%m%d')}.xlsx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=nome_arquivo,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def contar(inicio: date, fim: date) -> int:
    return RegistroPesoProduto.query.filter(
        RegistroPesoProduto.data >= inicio, RegistroPesoProduto.data <= fim
    ).count()


def linhas_combinadas(inicio: date, fim: date) -> list:
    registros = (
        RegistroPesoProduto.query.filter(
            RegistroPesoProduto.data >= inicio, RegistroPesoProduto.data <= fim
        )
        .order_by(RegistroPesoProduto.data.asc(), RegistroPesoProduto.horario.asc())
        .all()
    )
    return [
        {
            "tipo": "PAC 06-D - Peso do Produto",
            "icone": "⚖️",
            "data": r.data.strftime("%d/%m/%Y"),
            "horario": r.horario.strftime("%H:%M"),
            "resumo": f"{r.produto}: {r.total_pesagens} pesagem(ns), {r.total_nc} NC(s)",
            "conforme": r.conforme,
            "ordenacao": datetime.combine(r.data, r.horario),
        }
        for r in registros
    ]
