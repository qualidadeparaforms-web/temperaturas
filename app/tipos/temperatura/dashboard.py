from datetime import date, datetime, timedelta
from io import BytesIO

from flask import Blueprint, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.tipos.temperatura.models import ETAPAS, RegistroTemperatura

# Convenção: o blueprint de dashboard de cada tipo se chama
# "<slug>_dashboard", com url_prefix "/dashboard/<slug>" e uma rota
# "" (GET) chamada "index" — é assim que o painel principal
# (app/routes/dashboard.py) redireciona para o painel de cada tipo via
# url_for(f"{tipo.slug}_dashboard.index").
dashboard_bp = Blueprint("temperatura_dashboard", __name__, url_prefix="/dashboard/temperatura")

PERIODOS_VALIDOS = {"dia", "semana", "mes", "personalizado"}

COR_CABECALHO = "1F3864"
COR_FUNDO_FORA_PADRAO = "FFC7CE"
COR_FONTE_FORA_PADRAO = "C00000"


def periodo_para_datas(periodo: str, data_inicio: str | None = None, data_fim: str | None = None):
    """Converte o filtro de período em um intervalo (inicio, fim), inclusive."""
    hoje = date.today()

    if periodo == "dia":
        return hoje, hoje
    if periodo == "semana":
        return hoje - timedelta(days=6), hoje
    if periodo == "mes":
        return hoje - timedelta(days=29), hoje
    if periodo == "personalizado" and data_inicio and data_fim:
        try:
            return (
                datetime.strptime(data_inicio, "%Y-%m-%d").date(),
                datetime.strptime(data_fim, "%Y-%m-%d").date(),
            )
        except ValueError:
            pass

    # Padrão: última semana.
    return hoje - timedelta(days=6), hoje


def buscar_registros(periodo: str, etapa: str, data_inicio: str | None = None, data_fim: str | None = None):
    inicio, fim = periodo_para_datas(periodo, data_inicio, data_fim)
    query = RegistroTemperatura.query.filter(
        RegistroTemperatura.data >= inicio, RegistroTemperatura.data <= fim
    )
    if etapa and etapa != "todas":
        query = query.filter(RegistroTemperatura.etapa == etapa)

    registros = query.order_by(
        RegistroTemperatura.data.asc(), RegistroTemperatura.horario.asc()
    ).all()
    return registros, inicio, fim


def _parametros_filtro():
    return {
        "periodo": request.args.get("periodo", "semana"),
        "etapa": request.args.get("etapa", "todas"),
        "data_inicio": request.args.get("data_inicio"),
        "data_fim": request.args.get("data_fim"),
    }


@dashboard_bp.route("")
def index():
    from app.tipos import TIPOS_REGISTRO

    return render_template("tipos/temperatura/dashboard.html", etapas=ETAPAS, tipos=TIPOS_REGISTRO)


@dashboard_bp.route("/api")
def api_registros():
    filtros = _parametros_filtro()
    registros, inicio, fim = buscar_registros(**filtros)

    total = len(registros)
    fora_padrao = sum(1 for r in registros if not r.conforme)
    conformes = total - fora_padrao

    return jsonify(
        {
            "registros": [r.to_dict() for r in registros],
            "resumo": {
                "total": total,
                "conformes": conformes,
                "fora_padrao": fora_padrao,
                "percentual_conforme": round(conformes / total * 100, 1) if total else 100.0,
            },
            "periodo": {
                "inicio": inicio.strftime("%d/%m/%Y"),
                "fim": fim.strftime("%d/%m/%Y"),
            },
        }
    )


def adicionar_planilha(wb: Workbook, inicio: date, fim: date, filtros_extra: dict | None = None) -> None:
    """Acrescenta ao workbook uma aba com os registros de temperatura
    do período. Usada tanto pela exportação exclusiva deste tipo
    quanto pela exportação combinada de todos os tipos (uma aba por
    tipo) — por isso sempre cria uma aba nova (`wb.create_sheet`) em
    vez de assumir que é a única aba do arquivo.
    """
    filtros_extra = filtros_extra or {}
    etapa = filtros_extra.get("etapa", "todas")

    query = RegistroTemperatura.query.filter(
        RegistroTemperatura.data >= inicio, RegistroTemperatura.data <= fim
    )
    if etapa and etapa != "todas":
        query = query.filter(RegistroTemperatura.etapa == etapa)
    registros = query.order_by(RegistroTemperatura.data.asc(), RegistroTemperatura.horario.asc()).all()

    ws = wb.create_sheet("Temperatura")

    colunas = ["Dia", "Horário", "Etapa", "Temperatura", "Responsável", "Conforme"]
    ws.append(colunas)

    cabecalho_preenchimento = PatternFill(start_color=COR_CABECALHO, end_color=COR_CABECALHO, fill_type="solid")
    cabecalho_fonte = Font(color="FFFFFF", bold=True)
    for indice in range(1, len(colunas) + 1):
        celula = ws.cell(row=1, column=indice)
        celula.fill = cabecalho_preenchimento
        celula.font = cabecalho_fonte
        celula.alignment = Alignment(horizontal="center")

    preenchimento_fora_padrao = PatternFill(
        start_color=COR_FUNDO_FORA_PADRAO, end_color=COR_FUNDO_FORA_PADRAO, fill_type="solid"
    )
    fonte_fora_padrao = Font(color=COR_FONTE_FORA_PADRAO, bold=True)

    for registro in registros:
        ws.append(
            [
                registro.data.strftime("%d/%m/%Y"),
                registro.horario.strftime("%H:%M"),
                registro.etapa,
                registro.temperatura,
                registro.responsavel,
                "Sim" if registro.conforme else "Não",
            ]
        )
        if not registro.conforme:
            linha = ws.max_row
            for indice in range(1, len(colunas) + 1):
                celula = ws.cell(row=linha, column=indice)
                celula.fill = preenchimento_fora_padrao
                celula.font = fonte_fora_padrao

    larguras = [12, 10, 14, 14, 24, 12]
    for indice, largura in enumerate(larguras, start=1):
        ws.column_dimensions[chr(64 + indice)].width = largura


@dashboard_bp.route("/exportar")
def exportar():
    filtros = _parametros_filtro()
    inicio, fim = periodo_para_datas(filtros["periodo"], filtros["data_inicio"], filtros["data_fim"])

    wb = Workbook()
    del wb["Sheet"]  # adicionar_planilha sempre cria sua própria aba nomeada
    adicionar_planilha(wb, inicio, fim, {"etapa": filtros["etapa"]})

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nome_arquivo = (
        f"registros_temperatura_{inicio.strftime('%Y%m%d')}_{fim.strftime('%Y%m%d')}.xlsx"
    )
    return send_file(
        buffer,
        as_attachment=True,
        download_name=nome_arquivo,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def contar(inicio: date, fim: date) -> int:
    """Quantidade de registros no período — usado nos cartões da
    visão combinada do painel (quando há mais de um tipo de registro)."""
    return RegistroTemperatura.query.filter(
        RegistroTemperatura.data >= inicio, RegistroTemperatura.data <= fim
    ).count()


def linhas_combinadas(inicio: date, fim: date) -> list:
    """Lista padronizada para a tabela da visão combinada do painel
    (quando há mais de um tipo de registro)."""
    registros = (
        RegistroTemperatura.query.filter(
            RegistroTemperatura.data >= inicio, RegistroTemperatura.data <= fim
        )
        .order_by(RegistroTemperatura.data.asc(), RegistroTemperatura.horario.asc())
        .all()
    )
    return [
        {
            "tipo": "Temperatura de Processo",
            "icone": "🌡️",
            "data": r.data.strftime("%d/%m/%Y"),
            "horario": r.horario.strftime("%H:%M"),
            "resumo": f"{r.etapa}: {r.temperatura:.1f}°C",
            "conforme": r.conforme,
            "ordenacao": datetime.combine(r.data, r.horario),
        }
        for r in registros
    ]
